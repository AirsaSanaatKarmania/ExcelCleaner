#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
پاک‌سازی ارقام و تبدیل اکسل به «شکل خام» (برنامه ویندوز)

  ۱) ارقام فارسی/عربی  ←  انگلیسی            ۱۲۳  →  123
  ۲) متن‌های عددی با جداکننده  ←  عدد خام     5892-1041-0085-9016  →  5892104100859016
  ۳) تبدیل به شکل خام: فرمول ← مقدار، حذف قالب و فرمت سلول‌ها، باز کردن ادغام،
     حذف لینک/کامنت/قالب‌بندی شرطی/Validation/جدول و فیلتر
  • کاربر خودش سلول‌ها را انتخاب می‌کند (ماوس، ستون، ردیف یا آدرس مثل B2:B500)
    یا کل شیت / همه شیت‌ها را خام می‌کند.

نصب:   pip install openpyxl tksheet
اجرا:  py excel_cleaner.py
ساخت exe: build_exe.bat
"""
import re
import tkinter as tk
from datetime import date, datetime, time, timedelta
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.styles.cell_style import StyleArray
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.datetime import to_excel
from tksheet import Sheet

__version__ = "1.0.0"

FONT = ("Tahoma", 10)
FONT_B = ("Tahoma", 10, "bold")
GREEN = "#b7efc5"

# ───────────────────────── منطق تبدیل ─────────────────────────

DIGIT_MAP = {}
for _i, _ch in enumerate("۰۱۲۳۴۵۶۷۸۹"):      # فارسی
    DIGIT_MAP[ord(_ch)] = str(_i)
for _i, _ch in enumerate("٠١٢٣٤٥٦٧٨٩"):      # عربی
    DIGIT_MAP[ord(_ch)] = str(_i)
DIGIT_MAP[ord("٫")] = "."                      # ممیز فارسی
DIGIT_MAP[ord("٬")] = ","                      # جداکننده هزارگان فارسی

SEP_CHARS = {
    "space": r"\s\u200b-\u200f\u202a-\u202e\u2060\u00a0\ufeff",   # فاصله، نیم‌فاصله، نامرئی‌ها
    "dash": r"\-\u2010-\u2015\u2212_",                             # - ‐ – — − _
    "comma": r",،",
    "dot": r"\.",
    "slash": r"/\\",
}
_INVIS = r"\s\u200b-\u200f\u202a-\u202e\u2060\u00a0\ufeff"
TRIM_RE = re.compile(rf"^[{_INVIS}]+|[{_INVIS}]+$")
ASCII_DIGITS = re.compile(r"[0-9]+")
DATE_TYPES = (datetime, date, time, timedelta)


def build_sep_regex(flags):
    chars = "".join(SEP_CHARS[k] for k, on in flags.items() if on)
    return re.compile(f"[{chars}]") if chars else None


def convert_text(s, to_en, strip, sep_re, as_number):
    """رشته → مقدار جدید (str یا int)، یا None اگر تغییری لازم نیست."""
    t = s.translate(DIGIT_MAP) if to_en else s
    result = t
    if strip and sep_re is not None:
        stripped = sep_re.sub("", t)
        if stripped and re.fullmatch(r"\d+", stripped):          # کاملاً عددی بعد از حذف جداکننده‌ها
            result = stripped
            if (as_number and stripped.isascii() and len(stripped) <= 15
                    and (stripped[0] != "0" or len(stripped) == 1)):
                return int(stripped)
    return result if result != s else None


def parse_address(text, nrows, ncols):
    """'A2:A100, C:D, 5:7, F5' → [(r1,c1,r2,c2), ...]  (۰-مبنا، انتهای باز)"""
    text = text.translate(DIGIT_MAP).strip()
    boxes = []
    for part in filter(None, re.split(r"[,\s;،]+", text)):
        m = re.fullmatch(r"([A-Za-z]{1,3})(\d+):([A-Za-z]{1,3})(\d+)", part)
        if m:
            c1, r1, c2, r2 = m.groups()
            a, b = sorted((column_index_from_string(c1.upper()), column_index_from_string(c2.upper())))
            x, y = sorted((int(r1), int(r2)))
            boxes.append((x - 1, a - 1, y, b))
            continue
        m = re.fullmatch(r"([A-Za-z]{1,3}):([A-Za-z]{1,3})", part)
        if m:
            a, b = sorted(column_index_from_string(g.upper()) for g in m.groups())
            boxes.append((0, a - 1, nrows, b))
            continue
        m = re.fullmatch(r"(\d+):(\d+)", part)
        if m:
            x, y = sorted(int(g) for g in m.groups())
            boxes.append((x - 1, 0, y, ncols))
            continue
        m = re.fullmatch(r"([A-Za-z]{1,3})(\d+)", part)
        if m:
            c, r = column_index_from_string(m.group(1).upper()), int(m.group(2))
            boxes.append((r - 1, c - 1, r, c))
            continue
        m = re.fullmatch(r"([A-Za-z]{1,3})", part)
        if m:
            c = column_index_from_string(m.group(1).upper())
            boxes.append((0, c - 1, nrows, c))
            continue
        raise ValueError(f"آدرس نامعتبر: {part}")
    clipped = []
    for r1, c1, r2, c2 in boxes:
        r1, c1, r2, c2 = max(r1, 0), max(c1, 0), min(r2, nrows), min(c2, ncols)
        if r1 < r2 and c1 < c2:
            clipped.append((r1, c1, r2, c2))
    return clipped


def show_value(v):
    return "" if v is None else str(v)


# ───────────────────────── برنامه ─────────────────────────

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("پاک‌سازی ارقام و تبدیل اکسل به شکل خام")
        self.geometry("1020x900")
        self.minsize(860, 700)
        self.option_add("*Font", FONT)

        self.wb = None
        self.wb_vals = None                      # نسخه data_only برای مقدار فرمول‌ها (تنبل)
        self.path = None
        self.nrows = self.ncols = 0
        self.changed = {}                        # نام شیت → مجموعه (r, c) سلول‌های تغییر کرده

        # تنظیمات بخش ۳
        self.var_en = tk.BooleanVar(value=True)
        self.var_strip = tk.BooleanVar(value=True)
        self.var_out = tk.StringVar(value="text")
        self.seps = {k: tk.BooleanVar(value=(k in ("space", "dash"))) for k in SEP_CHARS}
        # تنظیمات بخش ۴ (شکل خام)
        self.var_scope = tk.StringVar(value="sheet")
        self.raw = {
            "formulas": tk.BooleanVar(value=True),
            "styles": tk.BooleanVar(value=True),
            "unmerge": tk.BooleanVar(value=True),
            "links": tk.BooleanVar(value=True),
            "sheetlevel": tk.BooleanVar(value=True),
            "trim": tk.BooleanVar(value=True),
            "text": tk.BooleanVar(value=True),
            "serial": tk.BooleanVar(value=False),
        }
        self.var_sheet = tk.StringVar()
        self.var_addr = tk.StringVar()
        self.var_file = tk.StringVar()

        self._build()

    # ───── رابط ─────
    def _build(self):
        pad = {"padx": 10, "pady": 3}

        top = ttk.Frame(self)
        top.pack(fill="x", **pad)
        ttk.Button(top, text="۱) انتخاب فایل اکسل...", command=self.open_file).pack(side="right")
        ttk.Entry(top, textvariable=self.var_file, state="readonly", justify="left").pack(
            side="right", fill="x", expand=True, padx=8)
        self.cmb = ttk.Combobox(top, textvariable=self.var_sheet, state="readonly", width=22)
        self.cmb.pack(side="left")
        self.cmb.bind("<<ComboboxSelected>>", lambda e: self.show_sheet(self.var_sheet.get()))
        ttk.Label(top, text="شیت:").pack(side="left", padx=(0, 4))

        mid = ttk.LabelFrame(self, text="۲) سلول‌های مورد نظر را انتخاب کنید")
        mid.pack(fill="both", expand=True, **pad)

        tools = ttk.Frame(mid)
        tools.pack(fill="x", padx=6, pady=(4, 2))
        ttk.Label(tools, text="آدرس (مثال: B2:B500, D:D, 5:8):").pack(side="right", padx=3)
        ent = ttk.Entry(tools, textvariable=self.var_addr, width=30, justify="left")
        ent.pack(side="right", padx=3)
        ent.bind("<Return>", lambda e: self.select_by_address())
        ttk.Button(tools, text="انتخاب با آدرس", command=self.select_by_address).pack(side="right", padx=3)
        ttk.Button(tools, text="پاک کردن انتخاب", command=self.clear_sel).pack(side="right", padx=3)
        ttk.Button(tools, text="انتخاب همه", command=self.select_all).pack(side="right", padx=3)

        self.sheet = Sheet(
            mid, data=[[]], table_font=("Tahoma", 10, "normal"), height=200,
            header_font=("Tahoma", 10, "bold"), index_font=("Tahoma", 10, "normal"))
        self.sheet.enable_bindings(
            "single_select", "toggle_select", "drag_select", "select_all", "column_select",
            "row_select", "column_width_resize", "double_click_column_resize", "arrowkeys", "copy")
        self.sheet.pack(fill="both", expand=True, padx=6, pady=4)
        self.sheet.extra_bindings("all_select_events", self._on_select)

        self.lbl_sel = ttk.Label(mid, text="", foreground="#555")
        self.lbl_sel.pack(anchor="e", padx=8, pady=(0, 4))
        self._on_select()

        # ── بخش ۳: ارقام و جداکننده‌ها
        opt = ttk.LabelFrame(self, text="۳) تبدیل ارقام و جداکننده‌ها (روی سلول‌های انتخاب‌شده)")
        opt.pack(fill="x", **pad)
        ttk.Checkbutton(opt, variable=self.var_en,
                        text="تبدیل ارقام فارسی و عربی به انگلیسی     ۱۲۳  ←  123").pack(anchor="e", padx=8, pady=1)
        ttk.Checkbutton(opt, variable=self.var_strip,
                        text="حذف جداکننده‌ها از متن‌های کاملاً عددی     5892-1041-0085-9016  ←  5892104100859016"
                        ).pack(anchor="e", padx=8, pady=1)
        row = ttk.Frame(opt)
        row.pack(anchor="e", padx=28, pady=1)
        ttk.Label(row, text="جداکننده‌هایی که حذف شوند:").pack(side="right", padx=6)
        for key, text in [("space", "فاصله و نیم‌فاصله"), ("dash", "خط تیره (-) و زیرخط"),
                          ("comma", "ویرگول"), ("dot", "نقطه"), ("slash", "اسلش (/)")]:
            ttk.Checkbutton(row, text=text, variable=self.seps[key]).pack(side="right", padx=6)
        outrow = ttk.Frame(opt)
        outrow.pack(anchor="e", padx=8, pady=(2, 2))
        ttk.Label(outrow, text="نتیجه به صورت:", font=FONT_B).pack(side="right")
        ttk.Radiobutton(outrow, text="متن (پیشنهادی — شماره کارت و کد پرسنلی)",
                        variable=self.var_out, value="text").pack(side="right", padx=8)
        ttk.Radiobutton(outrow, text="عدد (فقط تا ۱۵ رقم و بدون صفر اول)",
                        variable=self.var_out, value="number").pack(side="right", padx=8)
        ttk.Button(opt, text="✔  اعمال تبدیل ارقام/جداکننده روی سلول‌های انتخاب‌شده",
                   command=self.apply).pack(anchor="e", padx=8, pady=(2, 6), ipadx=8, ipady=2)

        # ── بخش ۴: شکل خام
        rawf = ttk.LabelFrame(self, text="۴) تبدیل به شکل خام (حذف قالب، فرمت و فرمول)")
        rawf.pack(fill="x", **pad)
        r0 = ttk.Frame(rawf)
        r0.pack(anchor="e", padx=8, pady=(3, 0))
        ttk.Label(r0, text="محدوده عملیات:", font=FONT_B).pack(side="right")
        for text, val in [("فقط سلول‌های انتخاب‌شده", "selection"), ("کل شیت جاری", "sheet"),
                          ("همه شیت‌ها", "all")]:
            ttk.Radiobutton(r0, text=text, variable=self.var_scope, value=val).pack(side="right", padx=8)

        def cb(parent, key, text):
            ttk.Checkbutton(parent, text=text, variable=self.raw[key]).pack(side="right", padx=8)

        r1 = ttk.Frame(rawf)
        r1.pack(anchor="e", padx=8, pady=1)
        cb(r1, "formulas", "فرمول‌ها ← مقدار")
        cb(r1, "styles", "حذف قالب سلول‌ها (فونت، رنگ، حاشیه، فرمت عدد، ...)")
        cb(r1, "unmerge", "باز کردن ادغام سلول‌ها")
        r2 = ttk.Frame(rawf)
        r2.pack(anchor="e", padx=8, pady=1)
        cb(r2, "links", "حذف لینک‌ها و کامنت‌ها")
        cb(r2, "sheetlevel", "حذف قالب‌بندی شرطی، Validation، جدول و فیلتر (فقط در «کل شیت/همه شیت‌ها»)")
        r3 = ttk.Frame(rawf)
        r3.pack(anchor="e", padx=8, pady=1)
        cb(r3, "trim", "حذف فاصله‌های اضافی ابتدا و انتهای متن")
        cb(r3, "text", "همزمان تنظیمات بخش ۳ (ارقام و جداکننده) هم اعمال شود")
        cb(r3, "serial", "تاریخ/ساعت به عدد سریال خام")
        ttk.Button(rawf, text="🧹  تبدیل به شکل خام", command=self.make_raw).pack(
            anchor="e", padx=8, pady=(3, 6), ipadx=8, ipady=2)

        btns = ttk.Frame(self)
        btns.pack(fill="x", **pad)
        ttk.Button(btns, text="💾  ۵) ذخیره فایل جدید...", command=self.save).pack(
            side="right", padx=4, ipadx=14, ipady=3)
        ttk.Button(btns, text="↺ لغو همه تغییرات", command=self.reload).pack(side="left", padx=4, ipady=3)

        self.log_box = tk.Text(self, height=5, state="disabled", wrap="word")
        self.log_box.pack(fill="x", padx=10, pady=(0, 8))

    def log(self, msg):
        self.log_box.configure(state="normal")
        self.log_box.insert("end", msg + "\n")
        self.log_box.see("end")
        self.log_box.configure(state="disabled")
        self.update_idletasks()

    # ───── فایل و شیت ─────
    def open_file(self):
        path = filedialog.askopenfilename(
            title="فایل اکسل را انتخاب کنید", filetypes=[("Excel", "*.xlsx *.xlsm")])
        if path:
            self.load(path)

    def load(self, path):
        p = Path(path)
        self.config(cursor="watch")
        self.update_idletasks()
        try:
            self.wb = load_workbook(p, keep_vba=p.suffix.lower() == ".xlsm")
        except Exception as e:
            self.config(cursor="")
            messagebox.showerror("خطا", f"فایل باز نشد:\n{e}\n\n(فقط xlsx و xlsm پشتیبانی می‌شود؛ فایل xls را ابتدا در اکسل Save As کنید.)")
            return
        self.config(cursor="")
        self.path = p
        self.wb_vals = None
        self.changed = {}
        self.var_file.set(str(p))
        names = self.wb.sheetnames
        self.cmb.configure(values=names)
        self.var_sheet.set(names[0])
        self.show_sheet(names[0])
        self.log(f"فایل بارگذاری شد: {p.name}  ({len(names)} شیت)")

    def reload(self):
        if self.path and messagebox.askyesno("تأیید", "همه تغییرات اعمال‌شده لغو و فایل دوباره بارگذاری شود؟"):
            self.load(self.path)

    def show_sheet(self, name):
        ws = self.wb[name]
        self.nrows, self.ncols = max(ws.max_row, 1), max(ws.max_column, 1)
        data = [[show_value(ws.cell(r, c).value) for c in range(1, self.ncols + 1)]
                for r in range(1, self.nrows + 1)]
        self.sheet.deselect("all", redraw=False)
        self.sheet.dehighlight_all(redraw=False)
        self.sheet.set_sheet_data(data, reset_highlights=True)
        for (r, c) in self.changed.get(name, ()):
            self.sheet.highlight_cells(row=r, column=c, bg=GREEN, fg="black", redraw=False)
        self.sheet.redraw()
        self._on_select()

    # ───── انتخاب ─────
    def selection_boxes(self):
        return [(max(r1, 0), max(c1, 0), min(r2, self.nrows), min(c2, self.ncols))
                for r1, c1, r2, c2 in self.sheet.get_all_selection_boxes()]

    def selected_cells(self):
        cells = set()
        for r1, c1, r2, c2 in self.selection_boxes():
            for r in range(r1, r2):
                for c in range(c1, c2):
                    cells.add((r, c))
        return cells

    def _on_select(self, *_):
        n = len(self.selected_cells()) if self.wb else 0
        self.lbl_sel.configure(text=f"سلول‌های انتخاب‌شده: {n:,}      |      کشیدن ماوس = محدوده  •  "
                                    "Ctrl+کلیک = چند محدوده  •  کلیک روی حرف ستون/شماره ردیف = کل آن")

    def select_all(self):
        if self.wb:
            self.sheet.select_all()
            self._on_select()

    def clear_sel(self):
        self.sheet.deselect("all")
        self._on_select()

    def select_by_address(self):
        if not self.wb:
            return
        try:
            boxes = parse_address(self.var_addr.get(), self.nrows, self.ncols)
        except ValueError as e:
            messagebox.showwarning("آدرس نامعتبر", str(e))
            return
        if not boxes:
            messagebox.showwarning("توجه", "آدرس واردشده خارج از محدوده داده‌های شیت است.")
            return
        self.sheet.deselect("all", redraw=False)
        for r1, c1, r2, c2 in boxes:
            self.sheet.create_selection_box(r1, c1, r2, c2, type_="cells")
        self.sheet.see(boxes[0][0], boxes[0][1])
        self.sheet.redraw()
        self._on_select()

    # ───── بخش ۳: تبدیل ارقام/جداکننده ─────
    def apply(self):
        if not self.wb:
            messagebox.showwarning("توجه", "ابتدا فایل اکسل را انتخاب کنید.")
            return
        cells = self.selected_cells()
        if not cells:
            messagebox.showwarning("توجه", "هیچ سلولی انتخاب نشده است.")
            return
        if not (self.var_en.get() or self.var_strip.get()):
            messagebox.showwarning("توجه", "حداقل یکی از گزینه‌های تبدیل را فعال کنید.")
            return

        sep_re = build_sep_regex({k: v.get() for k, v in self.seps.items()})
        as_number = self.var_out.get() == "number"
        name = self.var_sheet.get()
        ws = self.wb[name]
        done = self.changed.setdefault(name, set())

        n_changed = n_formula = n_merged = 0
        risky = []
        for (r, c) in sorted(cells):
            cell = ws.cell(r + 1, c + 1)
            if isinstance(cell, MergedCell):
                n_merged += 1
                continue
            v = cell.value
            if cell.data_type == "f":
                n_formula += 1
                continue
            if isinstance(v, str):
                new = convert_text(v, self.var_en.get(), self.var_strip.get(), sep_re, as_number)
                if new is None:
                    continue
                cell.value = new
                if isinstance(new, str):
                    cell.data_type = "s"
                if isinstance(new, str) and ASCII_DIGITS.fullmatch(new):
                    cell.number_format = "@"                 # متن بماند، با ویرایش بعدی هم عدد نشود
                elif isinstance(new, int) and cell.number_format == "@":
                    cell.number_format = "General"
                self.sheet.set_cell_data(r, c, show_value(new), redraw=False)
                self.sheet.highlight_cells(row=r, column=c, bg=GREEN, fg="black", redraw=False)
                done.add((r, c))
                n_changed += 1
            elif isinstance(v, (int, float)) and not isinstance(v, bool) and float(v).is_integer() and abs(v) >= 1e15:
                risky.append(f"{get_column_letter(c + 1)}{r + 1}")
        self.sheet.redraw()

        self.log(f"[{name}] از {len(cells):,} سلول انتخاب‌شده، {n_changed:,} سلول تغییر کرد (سبز).")
        if n_formula:
            self.log(f"   • {n_formula:,} سلول فرمول بود و دست نخورد.")
        if n_merged:
            self.log(f"   • {n_merged:,} سلول ادغام‌شده (Merge) نادیده گرفته شد.")
        self._warn_risky(risky)
        if n_changed:
            self.log("   برای نوشتن در فایل، دکمه «ذخیره فایل جدید» را بزنید.")

    def _warn_risky(self, risky):
        if risky:
            shown = ", ".join(risky[:8]) + (" ..." if len(risky) > 8 else "")
            self.log(f"   ⚠ {len(risky):,} سلول «عددی» با ۱۶ رقم یا بیشتر پیدا شد ({shown}). اکسل در نوع عدد فقط ۱۵ رقم "
                     "را نگه می‌دارد و رقم آخر این سلول‌ها احتمالاً از قبل خراب شده است.")

    # ───── بخش ۴: شکل خام ─────
    def ensure_values(self):
        if self.wb_vals is None:
            self.config(cursor="watch")
            self.update_idletasks()
            try:
                self.wb_vals = load_workbook(self.path, data_only=True)
            finally:
                self.config(cursor="")

    def make_raw(self):
        if not self.wb:
            messagebox.showwarning("توجه", "ابتدا فایل اکسل را انتخاب کنید.")
            return
        o = {k: v.get() for k, v in self.raw.items()}
        scope = self.var_scope.get()
        cur = self.var_sheet.get()
        if scope == "selection" and not self.selected_cells():
            messagebox.showwarning("توجه", "محدوده «فقط سلول‌های انتخاب‌شده» انتخاب شده ولی سلولی انتخاب نشده است.")
            return
        if o["text"] and not (self.var_en.get() or self.var_strip.get()):
            o["text"] = False
        if o["formulas"]:
            try:
                self.ensure_values()
            except Exception as e:
                messagebox.showerror("خطا", f"مقدار فرمول‌ها خوانده نشد:\n{e}")
                return

        sep_re = build_sep_regex({k: v.get() for k, v in self.seps.items()})
        as_number = self.var_out.get() == "number"
        sheets = list(self.wb.sheetnames) if scope == "all" else [cur]

        self.config(cursor="watch")
        self.update_idletasks()
        try:
            for name in sheets:
                ws = self.wb[name]
                vws = self.wb_vals[name] if (o["formulas"] and name in self.wb_vals.sheetnames) else None
                if scope == "selection":
                    boxes = self.selection_boxes()
                    stats = self._raw_sheet(ws, vws, o, sep_re, as_number, boxes=boxes)
                else:
                    stats = self._raw_sheet(ws, vws, o, sep_re, as_number, boxes=None)
                self._log_raw(name, stats, scope)
        except Exception as e:
            self.config(cursor="")
            messagebox.showerror("خطا", f"تبدیل کامل نشد:\n{e}")
            return
        self.config(cursor="")
        self.show_sheet(cur)
        self.log("   برای نوشتن در فایل، دکمه «ذخیره فایل جدید» را بزنید.")

    def _raw_sheet(self, ws, vws, o, sep_re, as_number, boxes):
        st = dict(formulas=0, nocache=0, styled=0, changed=0, unmerged=0, links=0, risky=[])
        whole = boxes is None

        # ۱) باز کردن ادغام
        if o["unmerge"]:
            for rng in list(ws.merged_cells.ranges):
                hit = whole or any(rng.min_row - 1 < r2 and rng.max_row > r1 and
                                   rng.min_col - 1 < c2 and rng.max_col > c1
                                   for r1, c1, r2, c2 in boxes)
                if hit:
                    ws.unmerge_cells(str(rng))
                    st["unmerged"] += 1

        # ۲) سلول‌های هدف
        if whole:
            targets = list(ws._cells.values())            # فقط سلول‌های موجود (بدون ساختن سلول خالی)
        else:
            coords = set()
            for r1, c1, r2, c2 in boxes:
                for r in range(r1, r2):
                    for c in range(c1, c2):
                        coords.add((r, c))
            targets = [ws.cell(r + 1, c + 1) for r, c in sorted(coords)]

        done = self.changed.setdefault(ws.title, set())
        removed_refs = set()
        for cell in targets:
            if isinstance(cell, MergedCell):
                continue
            changed = False
            had_style = cell.has_style

            # فرمول ← مقدار
            if cell.data_type == "f" and o["formulas"]:
                cv = vws.cell(cell.row, cell.column).value if vws is not None else None
                if cv is None:
                    st["nocache"] += 1
                else:
                    cell.value = cv
                    if isinstance(cv, str):
                        cell.data_type = "s"
                    st["formulas"] += 1
                    changed = True

            if cell.data_type != "f":
                v = cell.value
                if isinstance(v, str):
                    new = v
                    if o["trim"]:
                        new = TRIM_RE.sub("", new)
                    if o["text"]:
                        conv = convert_text(new, self.var_en.get(), self.var_strip.get(), sep_re, as_number)
                        if conv is not None:
                            new = conv
                    if new != v:
                        cell.value = new
                        if isinstance(new, str):
                            cell.data_type = "s"
                        changed = True
                elif isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
                    cell.value = int(v)
                elif isinstance(v, (int, float)) and not isinstance(v, bool) and float(v).is_integer() and abs(v) >= 1e15:
                    st["risky"].append(cell.coordinate)
                elif isinstance(v, DATE_TYPES) and o["serial"]:
                    try:
                        cell.value = to_excel(v)
                    except Exception:
                        pass

            # قالب سلول
            if o["styles"] and had_style:
                cell._style = StyleArray()
                st["styled"] += 1
            if o["styles"] or changed:
                val = cell.value
                if isinstance(val, DATE_TYPES):
                    cell.value = val                       # دوباره bind → فرمت تاریخ پیش‌فرض
                elif isinstance(val, str) and ASCII_DIGITS.fullmatch(val):
                    cell.number_format = "@"               # شناسه‌های عددی (کارت، کد) متن بمانند

            # لینک و کامنت
            if o["links"]:
                if cell._hyperlink is not None:
                    removed_refs.add(cell.coordinate)
                    cell._hyperlink = None
                    st["links"] += 1
                if cell.comment is not None:
                    cell.comment = None
                    st["links"] += 1

            if changed:
                st["changed"] += 1
                done.add((cell.row - 1, cell.column - 1))

        if o["links"] and removed_refs:
            ws._hyperlinks = [h for h in ws._hyperlinks if h.ref not in removed_refs]

        # ۳) سطح شیت (فقط وقتی کل شیت)
        if whole:
            if o["links"]:
                ws._hyperlinks = []
            if o["styles"]:
                for d in list(ws.column_dimensions.values()) + list(ws.row_dimensions.values()):
                    d._style = StyleArray()
            if o["sheetlevel"]:
                ws.conditional_formatting = ConditionalFormattingList()
                ws.data_validations.dataValidation = []
                for t in list(ws.tables):
                    del ws.tables[t]
                ws.auto_filter.ref = None
        return st

    def _log_raw(self, name, st, scope):
        parts = []
        if st["formulas"]:
            parts.append(f"{st['formulas']:,} فرمول ← مقدار")
        if st["styled"]:
            parts.append(f"قالب {st['styled']:,} سلول پاک شد")
        if st["changed"] - st["formulas"] > 0:
            parts.append(f"{st['changed'] - st['formulas']:,} سلول متن/ارقام اصلاح شد")
        if st["unmerged"]:
            parts.append(f"{st['unmerged']:,} ادغام باز شد")
        if st["links"]:
            parts.append(f"{st['links']:,} لینک/کامنت حذف شد")
        self.log(f"[{name}] شکل خام: " + ("، ".join(parts) if parts else "تغییری لازم نبود."))
        if st["nocache"]:
            self.log(f"   ⚠ {st['nocache']:,} فرمول مقدار ذخیره‌شده نداشت و به همین دلیل دست نخورد. "
                     "فایل را یک بار در اکسل باز و Save کنید و دوباره امتحان کنید.")
        if st["formulas"]:
            self.log("   (مقدار فرمول‌ها همان مقداری است که اکسل آخرین بار در فایل ذخیره کرده بود.)")
        self._warn_risky(st["risky"])

    # ───── ذخیره ─────
    def save(self):
        if not self.wb:
            messagebox.showwarning("توجه", "ابتدا فایل اکسل را انتخاب کنید.")
            return
        ext = self.path.suffix.lower()
        out = filedialog.asksaveasfilename(
            title="ذخیره فایل جدید", defaultextension=ext,
            initialfile=f"{self.path.stem}_fixed{ext}", filetypes=[("Excel", f"*{ext}")])
        if not out:
            return
        try:
            self.wb.save(out)
        except PermissionError:
            messagebox.showerror("خطا", "ذخیره نشد.\nاگر فایل در اکسل باز است آن را ببندید و دوباره تلاش کنید.")
            return
        except Exception as e:
            messagebox.showerror("خطا", str(e))
            return
        self.log(f"✔ ذخیره شد: {out}")
        messagebox.showinfo("تمام شد", f"فایل ذخیره شد:\n{out}")


if __name__ == "__main__":
    App().mainloop()
